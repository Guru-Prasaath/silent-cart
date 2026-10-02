"""Builds presentation/churn_presentation.pptx (12 slides + speaker notes) from the pipeline outputs."""
import json

import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from . import config as C

DARK, GREEN, MINT, TINT = "17332A", "2E7D5B", "BFE3CF", "EEF5F0"
ORANGE, BLUE, INK, MUTED, WHITE = "EB6834", "2A78D6", "1B1B1B", "5F6B66", "FFFFFF"
FONT = "Calibri"
W, H = 13.333, 7.5
FIG = C.FIGURES


def rgb(h):
    return RGBColor.from_string(h)


def money(x):
    return f"{'−' if x < 0 else '+' if x > 0 else ''}${abs(x) / 1000:,.1f}k"


def amount(x):
    """Unsigned dollar amount for prose ('loses $15.4k'), where the verb already carries the sign."""
    return f"${abs(x) / 1000:,.1f}k"


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.n = 0

    def slide(self, dark=False, title=None, kicker=None):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.n += 1
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb(DARK if dark else WHITE)
        if title:
            if kicker:
                self.text(s, kicker.upper(), 0.6, 0.38, 12, 0.3, size=11, color=GREEN if not dark else MINT, bold=True)
            self.text(s, title, 0.6, 0.62, 12.1, 0.95, size=27, bold=True, color=WHITE if dark else INK)
        if not dark and self.n > 1:
            self.text(s, f"Silent Cart  ·  FreshBasket loyalty churn  ·  {self.n}", 0.6, 7.05, 6, 0.3, size=9, color=MUTED)
        return s

    def text(self, s, text, x, y, w, h, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             italic=False, name=None):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, 0)
        paras = text if isinstance(text, list) else [text]
        for i, ptxt in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            runs = ptxt if isinstance(ptxt, list) else [(ptxt, {})]
            for rt, opt in runs:
                r = p.add_run()
                r.text = rt
                f = r.font
                f.name, f.size = FONT, Pt(opt.get("size", size))
                f.bold, f.italic = opt.get("bold", bold), italic
                f.color.rgb = rgb(opt.get("color", color))
            if len(paras) > 1:
                p.space_after = Pt(opt.get("space", 6) if isinstance(ptxt, list) else 6)
        if name:
            tb.name = name
        return tb

    def bullets(self, s, items, x, y, w, h, size=15, color=INK, gap=8):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, 0)
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(gap)
            head, rest = item if isinstance(item, tuple) else ("", item)
            for t, b, c in [("•  ", False, GREEN), (head, True, color), (rest, False, color)]:
                if t:
                    r = p.add_run()
                    r.text = t
                    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), b, rgb(c)
        return tb

    def card(self, s, x, y, w, h, fill=TINT, line=None):
        shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shp.adjustments[0] = 0.08
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
        if line:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(1)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    def stat(self, s, value, label, x, y, w, color=GREEN, size=40, label_color=MUTED, h=1.6, fill=TINT):
        self.card(s, x, y, w, h, fill=fill)
        self.text(s, value, x + 0.25, y + 0.18, w - 0.5, 0.75, size=size, bold=True, color=color)
        self.text(s, label, x + 0.25, y + 0.18 + size / 72 + 0.12, w - 0.5, h - size / 72 - 0.35, size=12, color=label_color)

    def badge(self, s, num, x, y, d=0.42, fill=GREEN):
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
        c.fill.solid()
        c.fill.fore_color.rgb = rgb(fill)
        c.line.fill.background()
        tf = c.text_frame
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, 0)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        r = p.add_run()
        r.text = str(num)
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(14), True, rgb(WHITE)

    def image(self, s, name, x, y, w, h):
        """Fit an image inside the box, preserving aspect ratio, centred."""
        path = FIG / name
        iw, ih = Image.open(path).size
        scale = min(w / iw, h / ih)
        pw, ph = iw * scale, ih * scale
        return s.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))

    def table(self, s, df, x, y, w, col_w, size=12, header_fill=DARK, highlight_row=None):
        rows, cols = df.shape[0] + 1, df.shape[1]
        t = s.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(0.4 * rows)).table
        for j, cw in enumerate(col_w):
            t.columns[j].width = Inches(cw)
        for i in range(rows):
            for j in range(cols):
                cell = t.cell(i, j)
                cell.margin_left = cell.margin_right = Inches(0.08)
                val = df.columns[j] if i == 0 else df.iat[i - 1, j]
                cell.text = str(val)
                para = cell.text_frame.paragraphs[0]
                para.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
                f = para.runs[0].font
                f.name, f.size = FONT, Pt(size)
                f.bold = i == 0 or (highlight_row is not None and i - 1 == highlight_row)
                f.color.rgb = rgb(WHITE if i == 0 else INK)
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(header_fill if i == 0 else (TINT if highlight_row is not None and i - 1 == highlight_row else WHITE))
        return t

    @staticmethod
    def notes(s, text):
        """Speaker notes: one idea per line (line breaks are kept in PowerPoint's notes pane)."""
        s.notes_slide.notes_text_frame.text = "\n".join(line.strip() for line in text.strip().splitlines())


def build() -> str:
    O = C.OUTPUTS
    k = json.loads((O / "key_metrics.json").read_text())
    comp = pd.read_csv(O / "model_comparison.csv")
    scen = pd.read_csv(O / "retention_scenarios.csv")
    ab = pd.read_csv(O / "ablation_study.csv")
    chk = pd.read_csv(O / "validation_design_checks.csv")
    seg = pd.read_csv(O / "segment_risk.csv")
    eng = pd.read_csv(O / "modelling_dataset.csv")
    dq = pd.read_csv(O / "data_quality_log.csv")
    tc = k["test_active"]
    A, B, C1, C2, D, E = (scen.iloc[i] for i in range(6))
    act = eng[eng.segment == "Active"]
    r = lambda m: act[m].churned.mean()
    rule = comp[(comp.model == "Recency rule") & (comp.split == "test") & (comp.population == "Active only")].iloc[0]
    tier = seg[(seg.population == "All eligible") & (seg.dimension == "tier")].set_index("segment_value").observed_churn
    abw = ab[ab.variant == "With recency"].reset_index(drop=True)
    sw = k["support_whatif"]
    from .report import shap_group_shares
    sh_sel, sh_ch = shap_group_shares(pd.read_csv(O / "shap_feature_importance.csv"))
    d = Deck()

    # 1 ── Title
    s = d.slide(dark=True)
    d.text(s, "SILENT CART  ·  FRESHBASKET LOYALTY CHURN PREDICTION & RETENTION", 0.8, 1.2, 11, 0.4, size=13, color=MINT, bold=True)
    d.text(s, "Most of our churn has already happened.\nHere is how to stop the rest.", 0.8, 1.8, 11.5, 2.2, size=44, bold=True, color=WHITE)
    d.text(s, "Who is at risk, why, what to do, and what it is worth: from 18 months of loyalty data on 2,600 members",
           0.8, 4.35, 11, 0.9, size=18, color=MINT)
    d.text(s, [[("Guru Prasaath D", {"bold": True, "color": WHITE, "size": 15}),
                ("   ·   Data Science  ·  Tailwyndz Propel Internship assessment", {"color": MINT})]], 0.8, 6.4, 11, 0.5, size=12)
    d.notes(s, """MAIN POINT: Most of the "churn" has already happened. I will show how to stop the rest.
    SAY:
    - Good morning. I am Guru Prasaath.
    - FreshBasket asked two questions: which loyalty members will stop shopping, and what should we do about it?
    - Today they send the same offer to everyone. That wastes money.
    - In the next 12 minutes I will show four things:
    1. The churn number hides two very different problems.
    2. Members show warning signs about 2 months before they leave.
    3. A model that catches 3 out of 4 leavers, on data it never saw.
    4. A targeted plan that makes money instead of losing it.
    NEXT: Let me start with the answer.""")

    # 2 ── Executive summary
    s = d.slide(title="The answer in one slide", kicker="Executive summary")
    stats = [
        (f"{k['share_churners_already_lapsed']:.0%}", "of 'churners' had already stopped buying before April. They need win-back, not prediction", ORANGE),
        (f"{tc['recall']:.0%}", f"of at-risk Active members caught at {tc['precision']:.0%} precision, on an unseen quarter (base rate {tc['churn_rate']:.1%})", GREEN),
        (f"{r(act.tickets_6m >= 3):.0%}", f"churn for Active members with 3+ support tickets, vs {r(act.tickets_6m == 0):.0%} with none", GREEN),
        (money(E.net_value), f"net from a next-best-action playbook vs {money(A.net_value)} for today's blanket coupon", GREEN),
    ]
    for i, (v, lab, col) in enumerate(stats):
        d.stat(s, v, lab, 0.6 + i * 3.08, 1.85, 2.88, color=col, size=40, h=2.3)
    d.card(s, 0.6, 4.5, 12.13, 1.95, fill=DARK)
    d.text(s, "RECOMMENDATION", 0.95, 4.75, 5, 0.3, size=11, bold=True, color=MINT)
    d.bullets(s, [("Split the KPI: ", "Lapsed (win-back) vs At-risk Active (retention), plus a 60-day no-purchase trigger"),
                  ("Score monthly, act selectively: ", "service call, personalised coupon or win-back, chosen per member by expected value"),
                  ("Prove it: ", "20% hold-out control group for one quarter before scaling")],
              0.95, 5.1, 11.5, 1.6, size=15, color=WHITE, gap=6)
    d.notes(s, f"""MAIN POINT: Four numbers tell the whole story.
    SAY:
    - {k['share_churners_already_lapsed']:.0%}: most members labelled "churned" had already stopped buying BEFORE April. They need a win-back offer, not a prediction.
    - {tc['recall']:.0%}: among members who are still shopping, the model finds {tc['recall']:.0%} of the ones who leave. When it flags someone, it is right {tc['precision']:.0%} of the time. Normally only about 1 in 10 leave.
    - {r(act.tickets_6m >= 3):.0%}: members with 3 or more support tickets leave at {r(act.tickets_6m >= 3):.0%}, compared with {r(act.tickets_6m == 0):.0%} for members with none. Service problems drive churn.
    - {amount(E.net_value)}: my targeted plan earns {amount(E.net_value)}. Today's coupon-for-everyone loses {amount(A.net_value)}.
    POINT AT: the dark box at the bottom. These are my 3 recommendations. I will come back to them at the end.
    NEXT: First, I checked whether the data could be trusted.""")

    # 3 ── Data quality
    s = d.slide(title="We fixed the data before trusting it, and found the label leaks", kicker="Data quality")
    items = [
        ("Hidden header row", "Every sheet has a title banner; a default read silently breaks columns"),
        ("Duplicates", "71 exact duplicates + 2 sign-flipped customer-month conflicts removed"),
        ("Corrupted spend", "55 negative, zero or 9-15x inflated values rebuilt from spend = transactions x basket"),
        ("Impossible zeros", "2,105 'categories / coupons' values in zero-transaction months set to 0"),
        ("Missing months", "953 member-months absent: treated as missing data, never as zero activity"),
        ("Leaky label columns", "MONTHS_OBSERVED, LAST_PURCHASE_DATE, INSUFFICIENT_HISTORY_FLAG use Apr-Jun data: excluded"),
    ]
    for i, (h_, t_) in enumerate(items):
        y = 1.75 + i * 0.84
        d.badge(s, i + 1, 0.6, y + 0.04)
        d.text(s, h_, 1.2, y, 5.2, 0.32, size=15, bold=True)
        d.text(s, t_, 1.2, y + 0.33, 5.3, 0.5, size=12, color=MUTED)
    d.image(s, "01_dq_spend_repair.png", 6.9, 1.6, 5.9, 4.0)
    d.text(s, [[("Leakage check: ", {"bold": True, "color": ORANGE}),
                (f"adding the leaky columns would inflate Active-member PR-AUC from {chk.iloc[1].pr_auc_active:.2f} to "
                 f"{chk.iloc[2].pr_auc_active:.2f}. An automated test also proves our features ignore all post-cut-off data.", {})]],
           6.9, 5.8, 5.9, 1.0, size=13)
    d.notes(s, f"""MAIN POINT: I cleaned the data first, and I found 3 columns that "cheat".
    SAY:
    - I ran {len(dq)} data checks. Every check is logged in a file.
    - Problems I fixed: a hidden header row, duplicate rows, 55 broken spend values, impossible zeros and missing months.
    - Broken spend: "spend = transactions x basket size" is true for 99.8% of rows. So I could rebuild the wrong values exactly, without deleting any data.
    - Most important: 3 columns in the label table use April-June data. That is the period we are trying to predict. It is like seeing the exam answers. This is called "leakage".
    - I removed them. With them, the score would look better than it really is: {chk.iloc[1].pr_auc_active:.2f} would become {chk.iloc[2].pr_auc_active:.2f}. An automatic test makes sure this cannot happen again.
    POINT AT: the orange dots on the chart. These are the broken spend values.
    NEXT: Once the data was clean, one chart changed my whole approach.""")

    # 4 ── Two churns
    s = d.slide(title="Two different churn problems hide inside one number", kicker="What churn really looks like")
    d.image(s, "03_recency_cliff_two_churns.png", 0.5, 1.6, 7.7, 5.3)
    d.stat(s, f"{k['n_lapsed']:,} · {k['churn_rate_lapsed']:.0%}", "LAPSED: no purchase in Jan-Mar 2024. They left before the window opened. Action: win-back campaign, identified by a rule",
           8.5, 1.75, 4.3, color=ORANGE, size=32, h=2.25)
    d.stat(s, f"{k['n_active']:,} · {k['churn_rate_active']:.1%}", "ACTIVE: bought in Jan-Mar 2024. A genuine early-warning problem. Action: model + targeted retention",
           8.5, 4.2, 4.3, color=GREEN, size=32, h=2.25)
    d.notes(s, f"""MAIN POINT: There are two different kinds of churn hidden in one number.
    SAY:
    - This chart shows churn by "months since the member's last purchase".
    - Bought last month: about 3% leave. Missed 1 month: {eng[eng.recency_months == 2].churned.mean():.0%}. Missed 2 months: {eng[eng.recency_months == 3].churned.mean():.0%}. Missed 3 or more: almost 100%.
    - So the official churn rate of {k['churn_rate_all']:.0%} mixes two groups:
    - LAPSED: {k['n_lapsed']} members who were already gone before April. A simple rule finds them.
    - ACTIVE: {k['n_active']:,} members still shopping. Only {k['churn_rate_active']:.1%} of them leave. This is where prediction really matters.
    - Any model looks brilliant on the lapsed group. So from now on I show results for ACTIVE members separately. That is the honest number.
    POINT AT: the blue bars (active) and then the orange bars (lapsed).
    NEXT: So what happens before an active member leaves?""")

    # 5 ── Fingerprint
    s = d.slide(title="Churn has a fingerprint, and a two-month window to act", kicker="Early-warning signals")
    d.image(s, "04_churn_fingerprint.png", 0.5, 1.55, 12.3, 3.95)
    tiles = [(f"{r(act.spend_change_pct < -0.4):.0%}", "churn if spend fell 40%+ vs prior quarter"),
             (f"{r(act.app_l3 < 0.5):.0%}", "churn with under 0.5 app sessions / month"),
             (f"{k['complaint_test']['churn_if_complaint']:.0%} vs {k['complaint_test']['churn_if_none']:.0%}", "churn with vs without a recent formal complaint (p < 0.001)")]
    for i, (v, lab) in enumerate(tiles):
        d.stat(s, v, lab, 0.6 + i * 4.1, 5.6, 3.9, color=ORANGE, size=28, h=1.3)
    d.notes(s, f"""MAIN POINT: Leaving has warning signs, about 2 months in advance.
    SAY:
    - I lined up every active member on their last purchase, and looked back 6 months.
    - Blue = members who stayed. Their behaviour is flat.
    - Orange = members who left. On average they buy less, use the app less and open fewer emails, and their support tickets jump.
    - This gives the business a 2-month window to act.
    - Three simple alarms that need no model:
    - spend drops by 40% or more: {r(act.spend_change_pct < -0.4):.0%} of these members leave.
    - app used less than once every 2 months: {r(act.app_l3 < 0.5):.0%} leave.
    - a formal complaint: {k['complaint_test']['churn_if_complaint']:.0%} leave, compared with {k['complaint_test']['churn_if_none']:.0%}.
    - All of these are statistically significant.
    POINT AT: the "Support tickets" panel. The orange line shoots up at the end.
    NEXT: Which signals matter most, and does membership tier protect us?""")

    # 6 ── Engagement & support, tier
    s = d.slide(title="Support experience and engagement matter. Tier does not.", kicker="What drives churn")
    d.image(s, "05_engagement_support_vs_churn.png", 0.5, 1.55, 8.2, 3.0)
    d.image(s, "12_shap_by_feature_group.png", 0.5, 4.6, 8.2, 2.35)
    d.bullets(s, [("Support: ", f"3+ tickets in 6 months means {r(act.tickets_6m >= 3):.0%} churn vs {r(act.tickets_6m == 0):.0%}. A complaint roughly doubles risk"),
                  ("Engagement: ", "app use and email opens are top-3 SHAP drivers in both models"),
                  ("Tier: ", f"Platinum churns like Silver ({tier['Platinum']:.0%} vs {tier['Silver']:.0%}); tier, age and tenure not significant (q > 0.6)"),
                  ("Signal split: ", f"behaviour + engagement = {sh_sel:.0%} (logistic) / {sh_ch:.0%} (boosting) of SHAP attribution")],
              9.0, 1.75, 3.8, 5.0, size=14, gap=12)
    d.notes(s, f"""MAIN POINT: Service problems and low engagement matter. Membership tier does not.
    SAY:
    - The brief asked about three things: tier, marketing engagement and support experience.
    - Support: YES. 3 or more tickets means {r(act.tickets_6m >= 3):.0%} churn, compared with {r(act.tickets_6m == 0):.0%}.
    - Engagement: YES. App use and email opens are among the strongest signals in both models.
    - Tier: NO. Platinum members leave as often as Silver members ({tier['Platinum']:.0%} vs {tier['Silver']:.0%}). The expensive perks are not creating loyalty.
    - The bar at the bottom shows where the model's signal comes from: mostly behaviour and engagement.
    POINT AT: the "Tier" bullet on the right.
    NEXT: Before trusting any model, I tested it the honest way.""")

    # 7 ── Validation design
    s = d.slide(title="Validated the way it will be used: train on the past, test on an unseen quarter", kicker="Method")
    phases = [("TRAIN", "Jul-23 & Oct-23 snapshots", "Labels: Jul-Dec 2023", MINT, INK),
              ("VALIDATE", "Jan-24 snapshot", "Model choice, calibration, threshold", "8CCBAA", INK),
              ("TEST · reported once", "Apr-24 snapshot", "Official Apr-Jun 2024 label", DARK, WHITE)]
    for i, (t1, t2, t3, fill, fc) in enumerate(phases):
        x = 0.6 + i * 4.1
        d.card(s, x, 1.75, 3.85, 1.45, fill=fill)
        d.text(s, t1, x + 0.25, 1.9, 3.4, 0.35, size=13, bold=True, color=fc)
        d.text(s, t2, x + 0.25, 2.27, 3.4, 0.35, size=17, bold=True, color=fc)
        d.text(s, t3, x + 0.25, 2.67, 3.4, 0.35, size=12, color=fc)
        if i < 2:
            arr = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 3.88), Inches(2.33), Inches(0.2), Inches(0.3))
            arr.fill.solid()
            arr.fill.fore_color.rgb = rgb(GREEN)
            arr.line.fill.background()
    d.bullets(s, [("Rolling-origin snapshots: ", "the churn definition is replayed at earlier cut-offs (99.8% agreement with the official label), giving honest out-of-time data. No random split"),
                  ("Leakage-safe features: ", "last 6 months only; recency, 3-vs-3-month trends, engagement, support, history, profile (42 features)"),
                  ("Imbalance: ", "class / sample weights, then Platt calibration so scores are real probabilities, then F1-optimal threshold on validation Active members"),
                  ("Four models: ", "recency rule (baseline), logistic regression, random forest, gradient boosting; tracked in MLflow"),
                  ("Scale-ready: ", "PySpark port of the feature pipeline matches pandas exactly on all 4 snapshots, ready for Databricks")],
              0.6, 3.5, 12.1, 2.4, size=14, gap=7)
    tiles = [(f"{chk.iloc[0].pr_auc_active:.2f}", "Out-of-time test, clean features (reported)", GREEN),
             (f"{chk.iloc[1].pr_auc_active:.2f}", "Random 5-fold CV: behaviour is stable over time", MUTED),
             (f"{chk.iloc[2].pr_auc_active:.2f}", "Random CV + leaky label columns: inflation avoided", ORANGE)]
    d.text(s, "Active-member PR-AUC under different validation set-ups (gradient boosting)", 0.6, 5.45, 12, 0.3, size=12, color=MUTED, italic=True)
    for i, (v, lab, col) in enumerate(tiles):
        d.stat(s, v, lab, 0.6 + i * 4.1, 5.8, 3.9, color=col, size=26, h=1.15)
    d.notes(s, """MAIN POINT: I tested the model the way it will really be used: learn from the past, predict the future.
    SAY:
    - A random split mixes past and future data. The results look better than they really are.
    - Instead I trained on July and October 2023, tuned on January 2024, and tested ONCE on the real April-June 2024 answers.
    - The model only sees the 6 months before each date. It never sees the future.
    - The scores are calibrated: "30% risk" really means about 30 in 100 such members leave. This matters for the money calculations later.
    - Engineering: every model run is tracked in MLflow, and the feature code also runs in Spark for Databricks. I tested that it gives identical results.
    POINT AT: the three boxes at the bottom. The honest test gives 0.81. The "cheating" columns would have given 0.87.
    NEXT: So how good is the model?""")

    # 8 ── Results
    s = d.slide(title="On Active members, machine learning roughly doubles what a recency rule can do", kicker="Model results · unseen Apr-Jun 2024 quarter")
    t = comp[(comp.split == "test") & (comp.population == "Active only")].copy()
    t["PR-AUC (95% CI)"] = [f"{a:.2f} ({lo:.2f}-{hi:.2f})" for a, lo, hi in zip(t.pr_auc, t.pr_auc_ci_low, t.pr_auc_ci_high)]
    tbl = pd.DataFrame({"Model": t.model, "PR-AUC (95% CI)": t["PR-AUC (95% CI)"],
                        "ROC-AUC": t.roc_auc.map("{:.2f}".format), "Precision": t.precision.map("{:.0%}".format),
                        "Recall": t.recall.map("{:.0%}".format), "F1": t.f1.map("{:.2f}".format)}).reset_index(drop=True)
    sel = int(tbl.index[tbl["Model"] == k["selected_model"]][0])
    d.table(s, tbl, 0.6, 1.75, 6.6, [1.85, 1.6, 0.9, 0.95, 0.7, 0.6], size=12, highlight_row=sel)
    d.image(s, "07_pr_curves.png", 7.4, 1.55, 5.5, 3.3)
    g = (k.get("gru") or {}).get("Active only")
    sel_txt = (f"{k['selected_model']} (best of the tabular models on validation). A GRU sequence model (PyTorch) scores "
               f"{g['pr_auc']:.2f}, likely better: shadow-test it next quarter" if g else
               f"{k['selected_model']} (best on validation). Gradient boosting is statistically tied")
    d.bullets(s, [("Selected: ", sel_txt),
                  ("At threshold {:.2f}: ".format(k['threshold']), f"catches {tc['recall']:.0%} of Active churners, {tc['precision']:.0%} precision; top 10% of scores = {tc['lift_top10pct']:.1f}x lift"),
                  ("Honest framing: ", "on all members every model scores PR-AUC > 0.94 because lapsed members are easy. That number would mislead")],
              0.6, 4.25, 6.6, 2.6, size=13, gap=9)
    d.image(s, "09_calibration.png", 7.4, 4.95, 5.5, 2.0)
    d.notes(s, f"""MAIN POINT: For active members, machine learning is about 2x better than a simple rule.
    SAY:
    - A simple rule ("has not bought recently") scores {rule.pr_auc:.2f}.
    - My machine-learning models score about 0.8. (1.0 is perfect. Random guessing would score about 0.10.)
    - Logistic regression won on the tuning data, so I chose it. It is simple and easy to explain. Gradient boosting is about the same.
    - At my chosen threshold the model flags {k['high_risk_active']} members, and it catches 3 out of 4 real leavers.
    - The calibration chart (bottom right) shows the risk percentages can be trusted.
    - I also tried a deep-learning model (a GRU). It scores slightly higher{f" ({g['pr_auc']:.2f})" if g else ""}. My plan: run it in the background for one quarter, and switch only if it stays better.
    POINT AT: the table. The bold row is the model I chose.
    NEXT: A score alone is not enough. People need to know WHY.""")

    # 9 ── Drivers (SHAP) per member
    s = d.slide(title="Every member gets a score and three reasons the CRM team can act on", kicker="Explainability (SHAP)")
    d.image(s, "11_shap_summary_active.png", 0.5, 1.5, 6.3, 5.45)
    preds = pd.read_csv(O / "churn_predictions.csv")
    ex = preds[(preds.segment == "Active") & (preds.predicted_label == 1)].iloc[[0, 3, 8]]
    d.text(s, "From outputs/churn_predictions.csv (high-risk Active members)", 7.1, 1.6, 5.7, 0.3, size=12, color=MUTED, italic=True)
    for i, row in enumerate(ex.itertuples()):
        y = 2.0 + i * 1.5
        d.card(s, 7.1, y, 5.7, 1.32, fill=TINT)
        d.text(s, f"Member {row.customer_id}", 7.35, y + 0.14, 3, 0.3, size=13, bold=True)
        d.text(s, f"{row.churn_probability:.0%} risk", 10.9, y + 0.12, 1.7, 0.35, size=16, bold=True, color=ORANGE, align=PP_ALIGN.RIGHT)
        d.text(s, row.top_3_drivers.replace("; ", "  ·  "), 7.35, y + 0.5, 5.25, 0.8, size=12, color=INK)
    d.text(s, "Drivers are SHAP contributions, grouped so the three reasons are always distinct (e.g. not three ways of saying 'stopped buying').",
           7.1, 6.5, 5.7, 0.5, size=11, color=MUTED)
    d.notes(s, """MAIN POINT: Every member gets a risk score, three reasons, and a ready-to-send message.
    SAY:
    - Left: what pushes risk up for the whole group: low purchases, low app and email use, and recent support tickets.
    - Right: what the store team sees for EACH member: the risk, plus 3 reasons in plain English.
    - One honest caution: these reasons explain the model. They do not prove cause. That is why I recommend a test before spending at scale.
    - The last step is automated too. An AI writes the message or call script from the same 3 reasons.
    - Safety checks block made-up promo codes, wrong apologies and any mention of tracking. If the AI fails a check, a safe template is used.
    POINT AT: one member card on the right.
    NEXT: Do we really need behaviour data, or is basic customer information enough?""")

    # 10 ── Ablation + segments
    s = d.slide(title="Behaviour predicts churn. Demographics alone are no better than chance.", kicker="Ablation & segments")
    d.image(s, "13_segment_risk_active.png", 0.5, 1.6, 7.1, 3.2)
    sa = seg[seg.population == "Active only"].set_index(["dimension", "segment_value"]).predicted_risk
    seg_tiles = [(f"{sa['city', 'Austin']:.0%} vs {sa['city', 'Nashville']:.0%}", "predicted risk, Austin vs Nashville (riskiest vs safest city)"),
                 (f"{sa['tier', 'Platinum']:.0%} vs {sa['tier', 'Silver']:.0%}", "Platinum vs Silver predicted risk; observed churn is flat across tiers"),
                 (f"{sa['signup_cohort', '2024']:.0%}", "predicted risk for 2024 joiners: a honeymoon period")]
    for i, (v, lab) in enumerate(seg_tiles):
        d.stat(s, v, lab, 0.6 + i * 2.37, 5.0, 2.22, color=GREEN, size=22, h=1.75)
    steps = [("Demographics only", abw.pr_auc_active[0]), ("+ Behavioural", abw.pr_auc_active[1]), ("+ Engagement", abw.pr_auc_active[2]),
             ("+ Support", abw.pr_auc_active[3]), ("All features", abw.pr_auc_active[4])]
    d.text(s, "PR-AUC on Active members (test quarter)", 7.9, 1.65, 4.9, 0.3, size=12, color=MUTED)
    for i, (lab, v) in enumerate(steps):
        y = 2.05 + i * 0.56
        d.text(s, lab, 7.9, y, 2.4, 0.4, size=14, bold=i in (1, 2))
        bar = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.2), Inches(y + 0.04), Inches(2.0 * v), Inches(0.3))
        bar.fill.solid()
        bar.fill.fore_color.rgb = rgb(GREEN if i else "B8C2BD")
        bar.line.fill.background()
        d.text(s, f"{v:.2f}", 10.3 + 2.0 * v, y, 0.6, 0.4, size=13, bold=True)
    d.bullets(s, [("Base rate is {:.1%}: ".format(k['churn_rate_active']), "demographics-only adds nothing"),
                  ("Engagement adds +{:.2f}, support +{:.2f} ".format(abw.pr_auc_active[2] - abw.pr_auc_active[1], abw.pr_auc_active[3] - abw.pr_auc_active[2]), "on top of behaviour"),
                  ("Segments ", "differ far less than behaviour: use them for logistics (which team, which channel), not for targeting")],
              7.9, 5.0, 4.9, 2.0, size=13, gap=8)
    d.notes(s, f"""MAIN POINT: Behaviour predicts churn. Age, city and tier alone do not.
    SAY:
    - I added the data groups one at a time (the bars on the right).
    - Demographics only: {abw.pr_auc_active[0]:.2f}. That is no better than guessing.
    - Add purchase behaviour: {abw.pr_auc_active[1]:.2f}. Add engagement: {abw.pr_auc_active[2]:.2f}. Add support: {abw.pr_auc_active[3]:.2f}.
    - Even without "months since last purchase", the model still scores about {ab[ab.variant != "With recency"].pr_auc_active.iloc[-1]:.1f}. So it finds real early warning signs.
    - Left chart, by city and tier: the differences are small. Austin and Columbus are the highest; members who joined in 2024 are the lowest.
    - So use segments to decide WHO acts (which store team), not WHO to target.
    NEXT: Now the most important part for the business: the money.""")

    # 11 ── Scenarios
    s = d.slide(title="Targeting turns retention from a cost into a return", kicker="Retention scenarios · this cohort, next quarter")
    d.image(s, "14_retention_scenarios.png", 0.5, 1.5, 8.6, 2.6)
    d.image(s, "15_scenario_sensitivity.png", 0.5, 4.2, 5.4, 2.75)
    d.stat(s, money(A.net_value), f"Blanket coupon to all {int(A.members_targeted):,} members: ${A.cost / 1000:,.1f}k spent", 9.4, 1.6, 3.4, color=ORANGE, size=34, h=1.45)
    d.stat(s, money(E.net_value), f"Next-best-action playbook: {int(E.members_targeted)} members, ${E.cost / 1000:,.1f}k spent, ~{E.churners_prevented:.0f} churners prevented",
           9.4, 3.2, 3.4, color=GREEN, size=34, h=1.6)
    d.bullets(s, [("Support outreach: ", f"call everyone with tickets = {money(C1.net_value)}; call only model-flagged = {money(C2.net_value)}"),
                  ("Model what-if: ", f"resolving tickets cuts their mean risk {sw['mean_risk_before']:.0%} → {sw['mean_risk_after']:.0%} (upper bound)"),
                  ("Robust: ", f"targeted coupon breaks even at {B.break_even_uplift:.0%} uplift (assumed 15%); blanket never does in range")],
              6.2, 5.0, 6.6, 2.0, size=13, gap=8)
    d.notes(s, f"""MAIN POINT: Targeting turns retention from a loss into a profit.
    SAY:
    - My assumptions, stated clearly: a $10 coupon lowers a member's churn risk by 15%. A $15 service call lowers it by 25%. A member's value is 6 months of their spending, at a 25% margin.
    - Today's coupon for everyone costs about ${A.cost / 1000:,.0f}k and loses {amount(A.net_value)}. Most people were staying anyway, or had already left.
    - My plan gives each member at most one action, and only when it pays back. About {E.churners_prevented:.0f} members are kept, for ${E.cost / 1000:,.1f}k. That is a {amount(E.net_value)} profit.
    - Calling everyone with support tickets loses money. Calling only the members the model flags makes money.
    - Targeting still breaks even if the offers work only about {max(B.break_even_uplift / C.COUPON_RELATIVE_UPLIFT, C2.break_even_uplift / C.OUTREACH_RELATIVE_UPLIFT):.0%} as well as I assumed. The coupon for everyone never breaks even.
    POINT AT: the red bar (today's approach), then the green bar (my plan).
    NEXT: So here is what I recommend.""")

    # 12 ── Recommendations
    s = d.slide(dark=True, title="Recommendations: a 90-day plan", kicker="What to do next")
    recs = [("Redefine the KPI", "Report Lapsed and At-risk Active separately; add a 60-day no-purchase trigger"),
            ("Run the playbook monthly", "Score all members; route each to service call, coupon, win-back or no action by expected value"),
            ("Fix the support loop", "Close tickets within the 2-month window; prioritise model-flagged members"),
            ("Free nudges first", "Spend drop of 40%+ or app use under 0.5/month triggers app push and preferred-category content"),
            ("Rethink tier perks", "Platinum ≠ loyal: pilot engagement-based rewards, not spend-only tiers"),
            ("Prove it", f"20% hold-out: ~{int(k['ab_test_15pct']['total_members'])} flagged members detect a 15% uplift; replace assumptions with measurements")]
    for i, (h_, t_) in enumerate(recs):
        col, row = i % 2, i // 2
        x, y = 0.6 + col * 6.15, 1.8 + row * 1.32
        d.card(s, x, y, 5.95, 1.15, fill="22473A")
        d.badge(s, i + 1, x + 0.22, y + 0.33, fill=GREEN)
        d.text(s, h_, x + 0.85, y + 0.14, 4.9, 0.35, size=16, bold=True, color=WHITE)
        d.text(s, t_, x + 0.85, y + 0.5, 4.9, 0.6, size=12, color=MINT)
    d.text(s, [[("Production path on the Tailwyndz stack:  ", {"bold": True, "color": WHITE}),
                ("Azure Data Factory → Delta (Databricks) → PySpark features → MLflow champion model → CRM queue + Dash Retention Console (Databricks Apps)  ·  monitor PR-AUC, calibration & PSI drift monthly", {"color": MINT})]],
           0.6, 5.95, 12.1, 0.8, size=13)
    d.notes(s, f"""MAIN POINT: Six actions for the next 90 days.
    SAY:
    1. Report the two kinds of churn separately, and add an alert after 60 days with no purchase.
    2. Score every member each month, and give each one the single best action.
    3. Fix support problems fast, inside the 2-month warning window.
    4. Try free nudges (app and email) before paid offers.
    5. Rethink tier perks. Platinum is not buying loyalty.
    6. Prove it before scaling: keep 20% of flagged members as a control group. About {int(k['ab_test_15pct']['total_members'])} flagged members are enough to measure a 15% effect.
    - Everything rebuilds with one command, and there is a live dashboard for the CRM team.
    OPTIONAL (about 60 seconds, only if time allows): open the live dashboard, click a high-risk member, click "Draft with AI", then move one slider in the simulator.
    - Thank you. I am happy to take questions.""")

    C.PRESENTATION.mkdir(exist_ok=True)
    path = C.PRESENTATION / "churn_presentation.pptx"
    d.prs.save(path)
    write_speaker_notes(d.prs)
    print(f"  deck -> {path.relative_to(C.ROOT)} ({d.n} slides)")
    return str(path)


def write_speaker_notes(prs, wpm: int = 130) -> str:
    """Export every slide's speaker notes to presentation/speaker_notes.md: an easy-to-read practice script."""
    label = {"MAIN POINT:": "**Main point:**", "SAY:": "**Say:**", "POINT AT:": "**Point at:**", "NEXT:": "**Next slide:**",
             "OPTIONAL": "**Optional**"}
    total, body = 0, []
    for i, s in enumerate(prs.slides, 1):
        texts = [sh.text_frame.text.replace("\n", " ") for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        title = next((t for t in texts if not t.isupper() and len(t) > 15), f"Slide {i}")
        note = s.notes_slide.notes_text_frame.text.strip()
        spoken = [ln for ln in note.splitlines() if not ln.startswith(("POINT AT:", "NEXT:", "OPTIONAL"))]
        words = sum(len(ln.split()) for ln in spoken)
        total += words
        body += ["---", "", f"## Slide {i}: {title}", "", f"*About {words / wpm * 60:.0f} seconds*", ""]
        for ln in note.splitlines():
            key = next((k for k in label if ln.startswith(k)), None)
            if key == "MAIN POINT:":
                body += [f"{label[key]} {ln[len(key):].strip()}", ""]
            elif key == "SAY:":
                body += [label[key], ""]
            elif key == "OPTIONAL":
                body += ["", f"> **Optional:** {ln[len(key):].strip(' :(').rstrip(')')}"]
            elif key:
                body += ["", f"{label[key]} {ln[len(key):].strip()}"]
            else:
                body.append(ln)
        body.append("")
    head = ["# Silent Cart: speaker notes (practice script)", "",
            f"**Total speaking time: about {total / wpm:.0f} minutes** at {wpm} words per minute. With pauses and slide changes "
            "this is 11-13 minutes, which fits the 10-15 minute slot.", "",
            "**How to use this script**",
            "- **Main point** is the one sentence the panel must remember from the slide. If you forget everything else, say this.",
            "- **Say** lists short ideas, not a speech to memorise. Say them in your own words.",
            "- **Point at** tells you where to look or point on the slide.",
            "- **Next slide** is the bridge sentence that leads into the next slide.",
            "- Slow down on **slide 4** (two kinds of churn) and **slide 11** (the money). Those are the moments the panel remembers.",
            "- Open https://silent-cart.onrender.com/ a few minutes before presenting so the live demo is awake.", ""]
    path = C.PRESENTATION / "speaker_notes.md"
    path.write_text("\n".join(head + body))
    return str(path)
